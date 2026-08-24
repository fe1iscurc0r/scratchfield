"""
文档解析模块 - 支持 PDF/DOCX/MD 文件解析
"""
import logging
import re
from pathlib import Path
from typing import List, Optional

logger = logging.getLogger(__name__)


class DocumentParser:
    """文档解析器 - 支持多种格式"""
    
    @staticmethod
    def parse(file_path: str) -> str | None:
        """解析文档，返回纯文本内容
        
        Args:
            file_path: 文件路径
            
        Returns:
            解析后的文本内容，失败返回 None
        """
        path = Path(file_path)
        if not path.exists():
            logger.error(f"文件不存在: {file_path}")
            return None
        
        suffix = path.suffix.lower()
        
        try:
            if suffix == '.pdf':
                return DocumentParser._parse_pdf(path)
            elif suffix in ('.docx', '.doc'):
                return DocumentParser._parse_docx(path)
            elif suffix in ('.md', '.markdown'):
                return DocumentParser._parse_markdown(path)
            elif suffix == '.txt':
                return DocumentParser._parse_text(path)
            else:
                logger.warning(f"不支持的文件格式: {suffix}")
                return None
        except Exception as e:
            logger.error(f"解析文件失败 {file_path}: {e}")
            return None
    
    @staticmethod
    def _parse_pdf(path: Path) -> str | None:
        """解析 PDF 文件"""
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            text_parts: list[str] = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            return '\n'.join(text_parts)
        except ImportError:
            logger.warning("pypdf 未安装，尝试使用 pdfminer...")
            try:
                from pdfminer.high_level import extract_text
                return extract_text(str(path))
            except ImportError:
                logger.error("PDF 解析库未安装，请执行: pip install pypdf")
                return None
    
    @staticmethod
    def _parse_docx(path: Path) -> str | None:
        """解析 Word 文档"""
        try:
            from docx import Document
            doc = Document(str(path))
            text_parts: list[str] = []
            for para in doc.paragraphs:
                if para.text.strip():
                    text_parts.append(para.text)
            # 同时提取表格内容
            for table in doc.tables:
                for row in table.rows:
                    row_text: list[str] = [cell.text for cell in row.cells]
                    text_parts.append(' | '.join(row_text))
            return '\n'.join(text_parts)
        except ImportError:
            logger.error("python-docx 未安装，请执行: pip install python-docx")
            return None
    
    @staticmethod
    def _parse_markdown(path: Path) -> str | None:
        """解析 Markdown 文件（去除标记语法 + 清洗 HTML 标签）"""
        try:
            import markdown as md_lib
            content = path.read_text(encoding='utf-8')
            html = md_lib.markdown(content)
            # 先移除<script>和<style>块，再去除其他HTML标签
            html = re.sub(r'<script.*?</script>', '', html, flags=re.DOTALL)
            html = re.sub(r'<style.*?</style>', '', html, flags=re.DOTALL)
            text = re.sub(r'<[^>]+>', '', html)
            return text.strip()
        except ImportError:
            logger.warning("markdown 库未安装，直接读取文本")
            return path.read_text(encoding='utf-8')
    
    @staticmethod
    def _parse_text(path: Path) -> str | None:
        """解析纯文本文件"""
        return path.read_text(encoding='utf-8')
    
    @staticmethod
    def get_supported_formats() -> list[str]:
        """获取支持的文件格式"""
        return ['.pdf', '.docx', '.doc', '.md', '.markdown', '.txt']
