"""卷163：法学判例导入管线测试。

覆盖：
- 案号抽取（含最高法指导性案例、全/半角括号、历史「字第」格式）
- 法院 / 审理程序 / 裁判日期 / 案由 / 当事人抽取
- docx / xlsx 解析（pdf 走缺依赖降级路径）
- 一行一案（xlsx 多案）语义
- 批量导入编排（单文件失败不炸整批、失败入队）
- flk 导入器（节流 ≥3s、严禁并发、响应解析）

严格使用**真实判决书格式** fixture（工单硬约束）。
"""

from __future__ import annotations

import pytest  # noqa: E402  (卷173 分层标注；与文件既有 import pytest 重复无害)

# 卷173 测试分层标注：smoke ⊂ core；未标注文件默认 full（pyproject.toml markers）
pytestmark = [pytest.mark.core]

import asyncio
import sys
import time
from pathlib import Path

import pytest

# 让 domains.* namespace 包可导入
REPO_ROOT = Path(__file__).resolve().parent.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from domains.law.importers import case_import as ci  # noqa: E402
from domains.law.importers import flk  # noqa: E402

# ===========================================================================
# fixture：真实判决书（覆盖最高法指导性案例等格式）
# ===========================================================================

SUPREME_GUIDING = """最高人民法院
民事判决书
（2018）最高法民终1234号
上诉人（原审原告）：中国某某集团有限公司，住所地北京市。
法定代表人：张某，董事长。
被上诉人（原审被告）：李四，男，1980年1月1日出生。
上诉人中国某某集团有限公司因与被上诉人李四买卖合同纠纷一案，不服北京市高级人民法院
（2017）京民初567号民事判决，向本院提起上诉。本院立案后，依法组成合议庭进行了审理。
本院认为，原审判决认定事实清楚，适用法律正确。
依照《中华人民共和国民事诉讼法》第一百七十条第一款第一项规定，判决如下：
驳回上诉，维持原判。
本判决为终审判决。
审判长　王五
审判员　赵六
二〇一八年九月二十七日
书记员　孙七
"""

BASIC_COURT_FIRST = """上海市浦东新区人民法院
民事判决书
（2021）沪0115民初8888号
原告：王五，男，1975年5月5日出生。
被告：赵六，女，1978年8月8日出生。
原告王五与被告赵六离婚纠纷一案，本院于2021年3月1日立案后，依法适用简易程序，
公开开庭进行了审理。
依照《中华人民共和国民法典》第一千零七十九条规定，判决如下：
准予原告王五与被告赵六离婚。
审判员　李法官
二〇二一年六月十五日
"""

RETRIAL = """中华人民共和国最高人民法院
民事裁定书
（2020）最高法民申5678号
再审申请人（一审原告、二审上诉人）：甲公司，住所地广东省深圳市。
被申请人（一审被告、二审被上诉人）：乙公司，住所地北京市。
再审申请人甲公司因与被申请人乙公司劳动争议一案，不服广东省高级人民法院
（2019）粤民终999号民事判决，向本院申请再审。
本院认为，甲公司的再审申请不符合《中华人民共和国民事诉讼法》第二百条规定的情形。
裁定如下：
驳回甲公司的再审申请。
审判长　周八
二〇二〇年十二月十日
"""

IP_COURT = """北京知识产权法院
民事判决书
（2019）京73民终456号
上诉人：某科技有限公司。
被上诉人：某网络科技有限公司。
上诉人某科技有限公司因与被上诉人某网络科技有限公司侵害计算机软件著作权纠纷一案，
不服北京市海淀区人民法院（2018）京0108民初1234号民事判决，向本院提起上诉。
本院于2019年4月1日立案后，依法组成合议庭进行了审理。
依照《中华人民共和国著作权法》第四十八条规定，判决如下：
驳回上诉，维持原判。
审判长　吴九
二〇一九年十一月二十日
"""

CRIMINAL = """广东省广州市中级人民法院
刑事判决书
（2018）粤01刑初234号
公诉机关：广东省广州市人民检察院。
被告人：陈某，男，1990年3月3日出生。
广东省广州市人民检察院指控被告人陈某犯诈骗罪，于2018年1月10日向本院提起公诉。
本院依法组成合议庭，公开开庭审理了本案。
依照《中华人民共和国刑法》第二百六十六条规定，判决如下：
被告人陈某犯诈骗罪，判处有期徒刑三年。
审判长　郑十
二〇一八年八月二十日
"""

EXEC_OBJECTION = """江苏省苏州市中级人民法院
执行裁定书
（2017）苏05执异12号
异议人：张某。
申请执行人：某银行。
被执行人：某公司。
本院在执行某银行与某公司借款合同纠纷一案中，异议人张某向本院提出书面异议。
本院受理后，依法组成合议庭进行了审查。
依照《中华人民共和国民事诉讼法》第二百二十五条规定，裁定如下：
驳回异议人张某的异议请求。
审判长　冯一
二〇一七年七月七日
"""


# ===========================================================================
# 案号抽取
# ===========================================================================

CASE_NO_CASES = [
    ("（2018）最高法民终1234号", "（2018）最高法民终1234号"),
    ("(2019)最高法民终1524号", "(2019)最高法民终1524号"),
    ("（2019）最高法民再152号", "（2019）最高法民再152号"),
    ("(2018)最高法行申3190号", "(2018)最高法行申3190号"),
    ("（2021）沪0115民初8888号", "（2021）沪0115民初8888号"),
    ("（2017）苏05执异12号", "（2017）苏05执异12号"),
    ("(2018)粤01刑初234号", "(2018)粤01刑初234号"),
    ("（2019）京73民终456号", "（2019）京73民终456号"),
]


@pytest.mark.parametrize("raw,expect", CASE_NO_CASES)
def test_extract_case_no_keeps_original_form(raw: str, expect: str) -> None:
    """案号须保留原括号形态（全/半角），因为它是 law 包 id_field。"""
    assert ci.extract_case_no(f"某某判决书\n{raw}\n正文") == expect


def test_extract_case_no_guiding_case() -> None:
    """最高法指导性案例格式（含「民终」多字类型代字）。"""
    assert ci.extract_case_no(SUPREME_GUIDING) == "（2018）最高法民终1234号"


# ===========================================================================
# 全字段抽取
# ===========================================================================

FULL_CASES = [
    (SUPREME_GUIDING, {
        "case_no": "（2018）最高法民终1234号", "court": "最高人民法院",
        "trial_level": "二审", "judgment_date": "2018-09-27",
        "cause_of_action": "买卖合同纠纷",
    }),
    (BASIC_COURT_FIRST, {
        "case_no": "（2021）沪0115民初8888号", "court": "上海市浦东新区人民法院",
        "trial_level": "一审", "judgment_date": "2021-06-15",
        "cause_of_action": "离婚纠纷",
    }),
    (RETRIAL, {
        "case_no": "（2020）最高法民申5678号", "court": "最高人民法院",
        "trial_level": "再审", "judgment_date": "2020-12-10",
        "cause_of_action": "劳动争议",
    }),
    (IP_COURT, {
        "case_no": "（2019）京73民终456号", "court": "北京知识产权法院",
        "trial_level": "二审", "judgment_date": "2019-11-20",
    }),
    (CRIMINAL, {
        "case_no": "（2018）粤01刑初234号", "court": "广东省广州市中级人民法院",
        "trial_level": "一审", "judgment_date": "2018-08-20",
    }),
    (EXEC_OBJECTION, {
        "case_no": "（2017）苏05执异12号", "court": "江苏省苏州市中级人民法院",
        "trial_level": "执行", "judgment_date": "2017-07-07",
    }),
]


@pytest.mark.parametrize("text,expect", FULL_CASES)
def test_extract_case_fields(text: str, expect: dict) -> None:
    """案号/法院/审理程序/裁判日期/案由抽取准确率在 fixture 上 100%。"""
    case = ci.extract_case(text)
    for field, want in expect.items():
        assert getattr(case, field) == want, f"{field}: {getattr(case, field)!r} != {want!r}"


def test_court_not_swallowed_by_verb() -> None:
    """正文动词（不服）不得被吞进法院名。"""
    assert ci.extract_court(IP_COURT) == "北京知识产权法院"


def test_retrial_not_misjudged_as_second_instance() -> None:
    """「申请再审」不得被判成二审。"""
    assert ci.extract_trial_level(RETRIAL) == "再审"


def test_parties_exclude_narrative_clause() -> None:
    """当事人抽取不得混入正文叙述句（含动词/罪名）。"""
    case = ci.extract_case(SUPREME_GUIDING)
    assert case.parties == ["上诉人中国某某集团有限公司", "被上诉人李四"]
    criminal = ci.extract_case(CRIMINAL)
    assert "被告人陈某" in criminal.parties
    assert all("犯" not in p for p in criminal.parties)


def test_cause_of_action_variants() -> None:
    """案由抽取覆盖多种真实表述。"""
    variants = [
        ("上诉人中国某某集团有限公司因与被上诉人李四买卖合同纠纷一案，不服", "买卖合同纠纷"),
        ("原告王五与被告赵六侵权责任纠纷一案，本院受理后", "侵权责任纠纷"),
        ("申请人甲公司与被申请人乙公司劳动争议一案，不服", "劳动争议"),
        ("案由：买卖合同纠纷", "买卖合同纠纷"),
        ("原告张三诉被告李四离婚纠纷一案", "离婚纠纷"),
        ("因房屋租赁合同纠纷一案，本院依法组成合议庭", "房屋租赁合同纠纷"),
    ]
    for text, expect in variants:
        assert ci.extract_cause_of_action(text) == expect, text


def test_judgment_date_prefers_explicit_label() -> None:
    """显式「裁判日期：」标签优先于正文日期（xlsx 导出列语义）。"""
    text = "案号：（2021）沪0115民初8888号　裁判日期：2021年6月15日\n全文：本院于2021年3月1日立案\n"
    assert ci.extract_judgment_date(text) == "2021-06-15"


def test_missing_fields_reported() -> None:
    """抽取缺失字段须如实记入 missing（不许口头声称完整）。"""
    case = ci.extract_case(CRIMINAL)
    # 刑事判决书无民事案由 → 应如实标记缺失
    assert "案由" in case.missing


# ===========================================================================
# 解析：docx / xlsx（pdf 走缺依赖降级）
# ===========================================================================

def _make_docx(path: Path, lines: list[str]) -> None:
    import docx

    d = docx.Document()
    for ln in lines:
        d.add_paragraph(ln)
    d.save(str(path))


def test_parse_docx(tmp_path: Path) -> None:
    p = tmp_path / "judgment.docx"
    _make_docx(p, SUPREME_GUIDING.splitlines())
    case = ci.extract_case(ci.parse_file(p))
    assert case.case_no == "（2018）最高法民终1234号"
    assert case.court == "最高人民法院"


def _make_xlsx(path: Path, rows: list[list[str]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(str(path))


def test_xlsx_one_row_one_case(tmp_path: Path) -> None:
    """xlsx 一行一案：一个文件产出多个判例，各自独立抽取。"""
    p = tmp_path / "cases.xlsx"
    _make_xlsx(p, [
        ["案号", "法院", "裁判日期", "案由", "全文"],
        ["（2021）沪0115民初8888号", "上海市浦东新区人民法院", "2021年6月15日", "离婚纠纷",
         "原告王五与被告赵六离婚纠纷一案。"],
        ["（2020）最高法民申5678号", "最高人民法院", "2020年12月10日", "劳动争议",
         "再审申请人甲公司因与被申请人乙公司劳动争议一案。"],
    ])
    cases = ci.iter_cases_from_xlsx(p)
    assert len(cases) == 2
    assert cases[0].case_no == "（2021）沪0115民初8888号"
    assert cases[0].judgment_date == "2021-06-15"
    assert cases[1].case_no == "（2020）最高法民申5678号"


def test_parse_unsupported_suffix(tmp_path: Path) -> None:
    p = tmp_path / "a.txt"
    p.write_text("hello", encoding="utf-8")
    with pytest.raises(ci.ParseError):
        ci.parse_file(p)


def test_parse_missing_file(tmp_path: Path) -> None:
    with pytest.raises(ci.ParseError):
        ci.parse_file(tmp_path / "nope.docx")


# ===========================================================================
# 批量导入编排
# ===========================================================================

def test_process_upload_docx(tmp_path: Path) -> None:
    p = tmp_path / "j.docx"
    _make_docx(p, SUPREME_GUIDING.splitlines())
    item = ci.process_upload("j.docx", p.read_bytes())
    assert item.ok and item.case is not None
    assert item.case.case_no == "（2018）最高法民终1234号"


def test_batch_single_failure_does_not_break_batch(tmp_path: Path) -> None:
    """单文件失败不炸整批：好的入库、坏的入队。"""
    good = tmp_path / "good.docx"
    _make_docx(good, BASIC_COURT_FIRST.splitlines())
    result = ci.process_batch([
        ("good.docx", good.read_bytes()),
        ("bad.pdf", b"%PDF-1.4 not a real pdf"),
        ("weird.xyz", b"nonsense"),
    ])
    assert result.total == 3
    assert result.imported >= 1
    assert result.queued >= 1


def test_batch_limit_files() -> None:
    files = [(f"f{i}.docx", b"x") for i in range(ci.MAX_FILES_PER_BATCH + 1)]
    with pytest.raises(ValueError):
        ci.process_batch(files)


def test_batch_limit_bytes() -> None:
    huge = b"0" * (ci.MAX_BATCH_BYTES + 1)
    with pytest.raises(ValueError):
        ci.process_batch([("huge.docx", huge)])


def test_to_paper_dict_maps_law_fields() -> None:
    case = ci.extract_case(SUPREME_GUIDING)
    d = case.to_paper_dict(source="pkulaw", license_note="个人订阅")
    assert d["case_no"] == "（2018）最高法民终1234号"
    assert d["source"] == "pkulaw"
    assert d["license_note"] == "个人订阅"
    # 卷164 起：tags 专供 AI 打标（JSON 数组），规则抽取的元数据落 notes
    assert "tags" not in d
    assert "买卖合同纠纷" in (d["notes"] or "")


# ===========================================================================
# flk 导入器
# ===========================================================================

def test_flk_interval_at_least_3s() -> None:
    assert flk.MIN_REQUEST_INTERVAL_SECONDS >= 3.0


def test_flk_no_concurrency_in_source() -> None:
    """严禁并发：源码不得出现 asyncio.gather。"""
    src = (REPO_ROOT / "domains" / "law" / "importers" / "flk.py").read_text(encoding="utf-8")
    assert "asyncio.gather" not in src


def test_flk_throttle_waits() -> None:
    async def _run() -> float:
        flk._last_request_ts = 0.0
        flk._mark_request()
        t0 = time.monotonic()
        await flk._respect_interval()
        return time.monotonic() - t0

    waited = asyncio.run(_run())
    assert waited >= 2.9


def test_flk_parse_response() -> None:
    payload = {
        "code": 200,
        "result": {"data": [
            {"id": "abc", "title": "中华人民共和国民法典", "office": "全国人大",
             "status": "有效", "publish": "2020-05-28"},
            {"title": "缺id被丢弃"},
        ]},
    }
    items = flk._parse_search_response(payload)
    assert len(items) == 1
    assert items[0].flk_id == "abc"
    assert items[0].to_paper_dict()["source"] == "flk"


def test_flk_sync_offline_returns_structure() -> None:
    """sync 入口返回结构正确（不触发真实网络：names=[] 时无抓取）。"""
    summary = asyncio.run(flk.sync_regulations([]))
    assert summary["status"] == "empty"
    assert summary["targets"] == []
    assert summary["results"] == []
